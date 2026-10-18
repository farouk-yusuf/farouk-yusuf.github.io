#!/usr/bin/env python3
"""Prepare the complete scoped Power BI model from the original council CSV.

This does not create a PBIX or execute DAX. The six tables preserve the distinct
monthly, matched-period, all-supplier comparison and candidate-pair grains.
"""
import argparse
import ast
from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timedelta
from decimal import Decimal
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from story_workbook import read_sheet
from bcc_source import source_model

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT/'projects/01_BCC_Payments'
OUT = PROJECT/'power-bi'
SOURCE = PROJECT/'outputs/BCC_Spend_and_Payment_Controls_Pack.xlsx'
AS_OF = date(2026,7,23)


def money(value, multiplier=1):
    return float((Decimal(str(value))*multiplier).quantize(Decimal('.01')))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-csv', type=Path, required=True,
                        help='Original council CSV. Kept outside the public repository.')
    args = parser.parse_args()
    model = source_model(args.source_csv, PROJECT)
    OUT.mkdir(exist_ok=True)
    tree = ast.parse((PROJECT/'scripts/03_finance_pack.py').read_text())
    corporate = next(ast.literal_eval(node.value) for node in tree.body if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='CORP' for t in node.targets))
    group = lambda name:'Corporate services (combined)' if name in corporate else name
    source_spend = read_sheet(SOURCE,'Spend by directorate')
    matched = [[name, *[money(value) for value in values]] for name,values in sorted(model['matched'].items())]
    report_groups = sorted(row[0] for row in matched)
    previous_monthly = defaultdict(lambda:[Decimal(0),0])
    with (PROJECT/'outputs/sql_results/01_monthly_spend_by_directorate.csv').open() as stream:
        for row in csv.DictReader(stream):
            key = (date.fromisoformat(row['month']+'-01'),group(row['directorate']))
            previous_monthly[key][0] += Decimal(row['spend'])
            previous_monthly[key][1] += int(row['payments'])
    assert set(previous_monthly) == set(model['monthly'])
    for key,(amount,count) in model['monthly'].items():
        assert count == previous_monthly[key][1]
        assert abs(amount-previous_monthly[key][0]) <= Decimal('.02'), key
    monthly_rows = [[month,report_group,money(total),count] for (month,report_group),(total,count) in sorted(model['monthly'].items())]
    month_rows = []
    for month in sorted({r[0] for r in monthly_rows}):
        fy = month.year if month.month>=4 else month.year-1
        month_rows.append([month,month.strftime('%b %Y'),month.year*100+month.month,f'{fy}/{str(fy+1)[2:]}',month<AS_OF.replace(day=1)])
    drivers = [[g,s,money(a),money(b),money(b-a)] for (g,s),(a,b) in sorted(model['suppliers'].items())]
    pairs = []
    for index,row in enumerate(read_sheet(SOURCE,'Possible duplicates'),1):
        first,last = [(datetime(1899,12,30)+timedelta(days=float(row[key]))).date() for key in ('first_date','second_date')]
        gap = int(row['days_apart'])
        assert (last-first).days==gap and 0<=gap<=3
        pairs.append([f'PAIR-{index:03d}',row['supplier'],group(row['directorate']),row['category'],money(row['amount']),first,last,gap,'Same day' if gap==0 else 'One to three days'])
    tables = {
        'DimDirectorate':(['ReportGroup'],[[g] for g in report_groups]),
        'DimMonth':(['MonthStart','MonthLabel','YearMonthSort','FinancialYear','IsCompleteMonth'],month_rows),
        'FactMatchedSpend':(['ReportGroup','PriorGBP','CurrentGBP','FY2025_26GBP'],matched),
        'FactMonthlySpend':(['MonthStart','ReportGroup','NetPaymentsGBP','PaymentLines'],monthly_rows),
        'FactSupplierDrivers':(['ReportGroup','Supplier','PriorGBP','CurrentGBP','ReportedMovementGBP'],drivers),
        'FactDuplicatePairs':(['PairID','Supplier','ReportGroup','Category','AmountGBP','FirstDate','SecondDate','DaysApart','ReviewPriority'],pairs),
    }
    # Every shared key must map to a real reporting group; never silently drop rows.
    for name,(headers,rows) in tables.items():
        if 'ReportGroup' in headers:
            assert {r[headers.index('ReportGroup')] for r in rows}<=set(report_groups),name
    prior = sum(Decimal(str(r[1])) for r in matched)
    current = sum(Decimal(str(r[2])) for r in matched)
    assert prior==Decimal('558638892.64')
    assert current==Decimal('677461112.00')
    # The original presentation workbook rounded each amount to £0.01m.
    for r in source_spend:
        values = [prior,current,sum(v[2] for v in model['matched'].values())] if r['report_group']=='Total' else model['matched'][r['report_group']]
        for value,key in zip(values,('1 Apr - 23 Jul 2025 £m','1 Apr - 23 Jul 2026 £m','FY2025/26 outturn £m')):
            assert abs(value-Decimal(r[key])*1000000) <= 5000
    signatures = Counter((r[1],Decimal(str(r[4])).quantize(Decimal('.01')),r[5],r[6]) for r in pairs)
    assert signatures == model['candidate_signatures'], 'Candidate evidence differs from the raw source.'
    assert len(pairs)==281 and sum(r[7]==0 for r in pairs)==73
    candidate_value = sum(Decimal(str(r[4])) for r in pairs)
    assert candidate_value==Decimal('4448515.38')
    assert sum(r[3] for r in monthly_rows)==247785
    assert all(date(2024,7,1)<=r[5]<=r[6]<=AS_OF and r[4]>=1000 for r in pairs)
    assert len(report_groups)==6 and len(month_rows)==25 and len(drivers)>18
    assert len({r[0] for r in pairs})==len(pairs)
    assert len({(r[0],r[1]) for r in drivers})==len(drivers)
    for g,a,b,_ in matched:
        subset = [r for r in drivers if r[0]==g]
        assert sum(Decimal(str(r[2])) for r in subset)==Decimal(str(a))
        assert sum(Decimal(str(r[3])) for r in subset)==Decimal(str(b))
    assert all(Decimal(str(r[4]))==Decimal(str(r[3]))-Decimal(str(r[2])) for r in drivers)
    checks = {
        'model_version':'2-complete-supplier-comparison',
        'source_csv':model['source'],
        'reference_workbook':str(SOURCE.relative_to(ROOT)),
        'source_monthly_csv':'projects/01_BCC_Payments/outputs/sql_results/01_monthly_spend_by_directorate.csv',
        'publication_end':AS_OF.isoformat(),
        'table_rows':{name:len(rows) for name,(_,rows) in tables.items()},
        'prior_gbp':str(prior),'current_gbp':str(current),
        'change_gbp':str(current-prior),'change_ratio':float((current-prior)/prior),
        'payment_lines':sum(r[3] for r in monthly_rows),
        'candidate_pairs':len(pairs),'candidate_value_gbp':str(candidate_value),
        'same_day_pairs':sum(r[7]==0 for r in pairs),
        'supplier_coverage':'All supplier display names/reporting groups present in either matched window; no Top N truncation. Private accounts pooled.',
        'previous_workbook_rounded_totals_gbp':{'prior':'558640000','current':'677460000','change':'118820000'},
        'validation_scope':'Recomputed exact monthly and matched-period aggregates from all 247,785 source rows. Supplier sums reconcile to each department to the penny. All 281 candidate signatures independently match the source. Does not execute Power BI/DAX.'
    }
    workbook = Workbook()
    notes = workbook.active
    notes.title = 'Read me'
    for row in [
        ['Birmingham Power BI data','Complete matched-period supplier comparison, exact source amounts; prepared data, not a PBIX'],
        ['Source','Original council CSV, 247,785 records; source checksum in data-checks.json'],
        ['Matched period','1 April–23 July in 2025 and 2026; exact amounts, not rounded workbook values'],
        ['Monthly scope','July 2024–July 2026; July 2026 is partial; do not use monthly totals for arbitrary matched dates'],
        ['Candidate scope','Full publication coverage; one amount per pair, not both amounts or confirmed loss'],
        ['Supplier scope','Every supplier display name/reporting group present in either matched window. No Top N truncation. Other dates are covered by the monthly aggregate, not the supplier comparison.'],
        ['Personal identifiers','Original account IDs, pseudonyms and payment descriptions omitted. Accounts linked to any redaction are pooled by reporting group. PairID is a snapshot review identifier.'],
        ['Refresh','This file is a prepared snapshot; refresh upstream analysis and rebuild to update it'],
        ['DimDirectorate grain','One row per combined reporting group'],
        ['DimMonth grain','One row per observed month; relates only to FactMonthlySpend'],
        ['FactMatchedSpend grain','One row per reporting group, fixed matched period and FY2025/26 total'],
        ['FactMonthlySpend grain','One row per month and reporting group; exact source amounts, net of credits'],
        ['FactSupplierDrivers grain','One standardised supplier display name per reporting group, with every matched-period account included. Private accounts pooled.'],
        ['FactDuplicatePairs grain','One refined candidate pair; no invoice-level confirmation'],
        ['Public data source','Birmingham City Observatory, Payments to Suppliers over £500 (OGL)'],
        ['Validation','Source totals, all supplier/group reconciliations and all candidate signatures checked. See data-checks.json; DAX and report behaviour must be checked in Desktop'],
        ['Existing report','Keep table/column names and relationships. See UPDATE_EXISTING_REPORT.md; replace the workbook and refresh.']
    ]:notes.append(row)
    notes.column_dimensions['A'].width=30
    notes.column_dimensions['B'].width=110
    for name,(headers,rows) in tables.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:sheet.append(row)
        table = Table(displayName=name,ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4',showRowStripes=True)
        sheet.add_table(table)
        sheet.freeze_panes='A2'
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='1B302B')
            cell.font=Font(color='FFFFFF',bold=True)
        for column in sheet.columns:
            key = column[0].value
            sheet.column_dimensions[column[0].column_letter].width=min(60,max(15,len(key)+2))
            for cell in column[1:]:
                if isinstance(cell.value,date):cell.number_format='yyyy-mm-dd'
                elif key.endswith('GBP'):cell.number_format='£#,##0.00;[Red]-£#,##0.00'
    filename = OUT/'BCC_PowerBI_Starter.xlsx'
    workbook.save(filename)
    # Reopen the actual delivered file and verify named-table coverage and row counts.
    loaded = load_workbook(filename,data_only=True)
    for name,(_,rows) in tables.items():
        assert list(loaded[name].tables)==[name] and loaded[name].max_row==len(rows)+1,name
    loaded.close()
    (OUT/'data-checks.json').write_text(json.dumps(checks,indent=2)+'\n')
    json.loads((OUT/'bcc-theme.json').read_text())
    with ZipFile(OUT/'BCC_PowerBI_Starter.zip','w',compression=ZIP_DEFLATED) as archive:
        for name in ('START_HERE.md','UPDATE_EXISTING_REPORT.md','BCC_PowerBI_Starter.xlsx','create-measures.dax','bcc-theme.json','data-checks.json'):
            archive.write(OUT/name,arcname=name)
    print(json.dumps(checks,indent=2))
    print('Created:',(OUT/'BCC_PowerBI_Starter.zip').relative_to(ROOT))


if __name__=='__main__':
    main()
