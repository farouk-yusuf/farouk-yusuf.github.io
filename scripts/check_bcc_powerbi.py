#!/usr/bin/env python3
"""Independently reconcile delivered supplier rows to the council CSV using SQLite."""
import argparse
import ast
import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import subprocess
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT/'projects/01_BCC_Payments'
OUT = PROJECT/'power-bi'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-csv', required=True, type=Path)
    args = parser.parse_args()
    workbook = load_workbook(OUT/'BCC_PowerBI_Starter.xlsx', data_only=True)
    baseline = load_workbook(BytesIO(subprocess.check_output(
        ['git', 'show', '8b69def:projects/01_BCC_Payments/power-bi/BCC_PowerBI_Starter.xlsx'], cwd=ROOT
    )), read_only=True)
    tables = {}
    for sheet in workbook.worksheets[1:]:
        assert list(sheet.tables) == [sheet.title]
        assert [c.value for c in sheet[1]] == list(next(baseline[sheet.title].values)), sheet.title
        assert not any(c.data_type == 'f' for row in sheet for c in row), 'Unexpected Excel formula'
        tables[sheet.title] = list(sheet.values)[1:]
    assert len(tables) == 6
    assert all(isinstance(r[0], datetime) and isinstance(r[4], bool) for r in tables['DimMonth'])
    assert sum(not r[4] for r in tables['DimMonth']) == 1
    checks = json.loads((OUT/'data-checks.json').read_text())
    assert {name: len(rows) for name, rows in tables.items()} == checks['table_rows']

    db = sqlite3.connect(':memory:')
    db.executescript('''
        CREATE TABLE payments(ordinal INTEGER, day TEXT, department TEXT,
                              beneficiary TEXT, pence INTEGER, account TEXT,
                              redacted INTEGER);
        CREATE TABLE mapping(department TEXT PRIMARY KEY, report_group TEXT);
    ''')
    with args.source_csv.open(encoding='utf-8-sig', newline='') as stream:
        dialect = csv.Sniffer().sniff(stream.read(8192), delimiters=',;\t')
        stream.seek(0)
        def records():
            for ordinal, row in enumerate(csv.DictReader(stream, dialect=dialect)):
                redacted = ''.join(row['beneficiary'].upper().replace('-', '').split()).startswith('REDACTEDPERSONALDATA')
                yield ordinal, row['date'], row['department'], row['beneficiary'], int(Decimal(row['amount'])*100), row['supplier_number'], int(redacted)
        db.executemany('INSERT INTO payments VALUES(?,?,?,?,?,?,?)', records())
    tree = ast.parse((PROJECT/'scripts/01_clean.py').read_text())
    mapping = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'dir_map' for t in n.targets))
    corporate = {'Council Management', 'Finance', 'HR, OD & Transformation', 'Strategy, Insight & Governance'}
    db.executemany('INSERT INTO mapping VALUES(?,?)', [(k, 'Corporate services (combined)' if v in corporate else v) for k,v in mapping.items()])
    assert db.execute('SELECT COUNT(*), SUM(pence) FROM payments').fetchone() == (247785, 387385201219)
    # Independent integer arithmetic and SQL grouping check the full exported key set,
    # not merely its grand total or the largest suppliers.
    result = db.execute('''
        WITH protected AS (
            SELECT DISTINCT account FROM payments WHERE redacted=1
        ), spellings AS (
            SELECT account, beneficiary, COUNT(*) AS n, MIN(ordinal) AS first_seen
            FROM payments WHERE account NOT IN (SELECT account FROM protected)
            GROUP BY account, beneficiary
        ), ranked AS (
            SELECT account, beneficiary,
              ROW_NUMBER() OVER(PARTITION BY account ORDER BY n DESC, first_seen) AS position
            FROM spellings
        )
        SELECT m.report_group,
          COALESCE(n.beneficiary, 'Individuals (names withheld)') AS supplier,
          SUM(CASE WHEN p.day BETWEEN '2025-04-01' AND '2025-07-23' THEN p.pence ELSE 0 END),
          SUM(CASE WHEN p.day BETWEEN '2026-04-01' AND '2026-07-23' THEN p.pence ELSE 0 END)
        FROM payments p JOIN mapping m ON p.department=m.department
        LEFT JOIN ranked n ON p.account=n.account AND n.position=1
        WHERE p.day BETWEEN '2025-04-01' AND '2025-07-23'
           OR p.day BETWEEN '2026-04-01' AND '2026-07-23'
        GROUP BY m.report_group, supplier
    ''').fetchall()
    actual = {(r[0], r[1]): tuple(int((Decimal(str(v))*100).quantize(Decimal(1))) for v in r[2:]) for r in tables['FactSupplierDrivers']}
    expected = {(group, supplier): (prior, current, current-prior) for group,supplier,prior,current in result}
    assert actual == expected, 'Delivered supplier rows differ from the independent source aggregation.'
    assert sum(r[0] == 'City Operations' for r in tables['FactSupplierDrivers']) == 738
    for group, prior, current, _ in tables['FactMatchedSpend']:
        rows = [values for (g,_),values in actual.items() if g == group]
        assert sum(r[0] for r in rows) == int(Decimal(str(prior))*100)
        assert sum(r[1] for r in rows) == int(Decimal(str(current))*100)
    expected_files = {'START_HERE.md','UPDATE_EXISTING_REPORT.md','BCC_PowerBI_Starter.xlsx','create-measures.dax','bcc-theme.json','data-checks.json'}
    with ZipFile(OUT/'BCC_PowerBI_Starter.zip') as archive:
        assert archive.testzip() is None and set(archive.namelist()) == expected_files
        assert all(archive.read(name) == (OUT/name).read_bytes() for name in expected_files)
    db.close()
    baseline.close()
    workbook.close()
    print(f'Passed: unchanged six-table schema; {len(actual):,} supplier/group rows independently reconciled to the penny; 738 City Operations rows; date types; complete ZIP contents.')


if __name__ == '__main__':
    main()
